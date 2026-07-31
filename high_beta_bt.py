import os
import math
import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import time
BASE = Path(os.getcwd())
TOP20_DIR = BASE / "Top 20 High Beta stock"
DAYWISE_DIR = BASE / "daywise stocks"
OUTPUT_FILE = BASE / "NIFTY_LargeMidcap_250_high_beta_tracker.xlsx"
BROKERAGE_PCT = 0.0001
STARTING_CAPITAL = 1000000.0

#styling
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="2F5496")
DATA_FONT = Font(name="Arial", size=10)
SMALL_FONT = Font(name="Arial", size=8)
BUY_FILL = PatternFill("solid", fgColor="C6EFCE")
SELL_FILL = PatternFill("solid", fgColor="FFC7CE")
HOLD_FILL = PatternFill("solid", fgColor="DDEBF7")
YELLOW_FILL = PatternFill("solid", fgColor="FFFF00")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
CASH_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FONT = Font(name="Arial", size=11, color="0000FF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
PCT_FMT = "0.00%"
NUM_FMT = "#,##0.00"

start_time = time.time()
parquet_lookup = {}
for file_path in DAYWISE_DIR.iterdir():
    if file_path.is_file() and file_path.suffix == ".parquet":
        parquet_lookup[file_path.stem] = file_path

stock_cache = {}

def load_stock_data(stock_name):
    sym = stock_name.upper()
    if sym in stock_cache:
        return stock_cache[sym]

    path = parquet_lookup.get(sym)
    if path is None:
        stock_cache[sym] = None
        return None

    df = pd.read_parquet(path)
    df["date_time"] = pd.to_datetime(df["date_time"]).dt.date
    df = df.dropna(subset=["date_time"]).sort_values("date_time").reset_index(drop=True)
    df = df.set_index("date_time")
    stock_cache[sym] = df
    return df

def get_close_of_date(stock_name, date):
    df = load_stock_data(stock_name)
    if df is None:
        return None

    try:
        idx = df.index.get_indexer([date], method="pad")[0]
        if idx == -1:
            return None
        return float(df.iloc[idx]["close"])
    except Exception:
        return None

def get_all_trading_dates(start, end):
    df = load_stock_data("RELIANCE")
    if df is None:
        df = load_stock_data("HDFCBANK")
    if df is None:
        return []

    mask = (df.index >= start) & (df.index < end)
    return sorted(df.index[mask].tolist())


def calc_brokerage(value):
    return round(value * BROKERAGE_PCT, 2)


top20_files = sorted(TOP20_DIR.glob("top20_high_beta_*.csv"))
rebalance_dates = []
rebalance_data = {}

for file_path in top20_files:
    date_str = file_path.stem.split("_")[-1]
    rebalance_date = pd.to_datetime(date_str).date()
    rebalance_df = pd.read_csv(file_path)
    rebalance_dates.append(rebalance_date)
    rebalance_data[rebalance_date] = rebalance_df

print(f"Loaded {len(rebalance_dates)} rebalance dates")


cash = STARTING_CAPITAL
portfolio = {}
trade_log = []
daily_records = []
all_periods = []


def log_trade(date, symbol, action, qty, price, brokerage, cash_before, cash_after, notes=""):
    trade_log.append(
        {
            "date": date,
            "symbol": symbol,
            "action": action,
            "qty": qty,
            "price": round(price, 2),
            "value": round(qty * price, 2),
            "brokerage": round(brokerage, 2),
            "cash_before": round(cash_before, 2),
            "cash_after": round(cash_after, 2),
            "notes": notes,
        }
    )


for index, rebalance_date in enumerate(rebalance_dates):
    rebalance_date_str = rebalance_date.strftime("%Y-%m-%d")
    top20_df = rebalance_data[rebalance_date]
    top20_symbols = top20_df["Symbol"].str.strip().tolist()
    beta_lookup = dict(zip(top20_df["Symbol"].str.strip(), top20_df["Beta"]))

    if index + 1 < len(rebalance_dates):
        next_rebalance_date = rebalance_dates[index + 1]
    else:
        next_rebalance_date = rebalance_date + pd.Timedelta(days=90)

    current_holdings = set(portfolio.keys())
    target_holdings = set(top20_symbols)

    holds = current_holdings.intersection(target_holdings)
    sells = current_holdings.difference(target_holdings)
    buys = target_holdings.difference(current_holdings)

    for symbol in sorted(holds):
        price = get_close_of_date(symbol, rebalance_date)
        if price and symbol in portfolio:
            portfolio[symbol]["last_price"] = price
            pnl_pct = (price - portfolio[symbol]["entry_price"]) / portfolio[symbol]["entry_price"]
            log_trade(rebalance_date, symbol, "HOLD", portfolio[symbol]["qty"], price,
                      0, cash, cash, f"Beta: {beta_lookup.get(symbol, 0):.4f}, PnL%: {pnl_pct:.2%}")

    sell_proceeds = 0.0
    for symbol in sorted(sells):
        if symbol not in portfolio:
            continue

        price = get_close_of_date(symbol, rebalance_date)
        if price is None:
            continue

        qty = portfolio[symbol]["qty"]
        entry_price = portfolio[symbol]["entry_price"]
        trade_value = qty * price
        brokerage = calc_brokerage(trade_value)
        proceeds = trade_value - brokerage

        pnl_pct = (price - entry_price) / entry_price if entry_price > 0 else 0
        pnl_abs = (price - entry_price) * qty - brokerage

        cash_before = cash
        cash += proceeds
        sell_proceeds += proceeds
        log_trade(rebalance_date, symbol, "SELL", qty, price, brokerage, cash_before, cash,
                    f"Rebalance exit PnL:{pnl_pct:+.2%} (Rs {pnl_abs:+,.0f})")
        del portfolio[symbol]

    buy_list = sorted(buys)
    if buy_list:
        cash_per_stock = cash / len(buy_list)

        for symbol in buy_list:
            price = get_close_of_date(symbol, rebalance_date)
            if price is None or price <= 0:
                continue

            beta = beta_lookup.get(symbol, 0)
            qty = math.floor(cash_per_stock / (price * (1 + BROKERAGE_PCT)))
            if qty <= 0:
                continue

            trade_value = qty * price
            brokerage = calc_brokerage(trade_value)
            total_cost = trade_value + brokerage

            if total_cost > cash + 0.01:
                qty -= 1
                if qty <= 0:
                    continue
                trade_value = qty * price
                brokerage = calc_brokerage(trade_value)
                total_cost = trade_value + brokerage

            cash_before = cash
            cash -= total_cost
            portfolio[symbol] = {
                "qty": qty,
                "entry_price": price,
                "entry_date": rebalance_date,
                "beta": beta,
                "last_price": price
            }
            log_trade(rebalance_date, symbol, "BUY", qty, price,
                        brokerage, cash_before, cash, f"Rebalance entry Beta:{beta:.4f}")

    trading_dates = get_all_trading_dates(rebalance_date, next_rebalance_date)
    for trading_date in trading_dates:
        for symbol in list(portfolio.keys()):
            stock_df = load_stock_data(symbol)
            if stock_df is None or trading_date not in stock_df.index:
                continue

            close_today = stock_df.loc[trading_date, "close"]
            if isinstance(close_today, pd.Series):
                close_today = float(close_today.iloc[-1])
            else:
                close_today = float(close_today)
            portfolio[symbol]["last_price"] = close_today

        stock_value = sum(holding["qty"] * holding["last_price"] for holding in portfolio.values())
        total_value = stock_value + cash
        daily_records.append(
            {
                "date": trading_date,
                "stock_value": stock_value,
                "cash": cash,
                "total_value": total_value,
                "num_stocks": len(portfolio),
            }
        )

    all_periods.append(
        {
            "rb_date": rebalance_date_str,
            "top20": top20_symbols,
            "buys": sorted(buys),
            "sells": sorted(sells),
            "holds": sorted(holds),
            "cash_at_end": cash,
        }
    )

    stock_value = sum(holding["qty"] * holding["last_price"] for holding in portfolio.values())
    print(
        f"  {rebalance_date_str} | B:{len(buys)} S:{len(sells)} H:{len(holds)} | "
        f"Portfolio: ₹ {stock_value + cash:,.0f} (Cash: Rs {cash:,.0f})"
    )

print("\nSimulation complete. Building Excel...")

daily_df = pd.DataFrame(daily_records)
if not daily_df.empty:
    daily_df = daily_df.sort_values("date").drop_duplicates(subset="date", keep="last").reset_index(drop=True)
    daily_df["daily_return"] = daily_df["total_value"].pct_change()
    daily_df["peak"] = daily_df["total_value"].cummax()
    daily_df["drawdown"] = (daily_df["total_value"] - daily_df["peak"]) / daily_df["peak"]

wb = Workbook()

def style_header_row(ws, row, max_col):
    for column in range(1, max_col + 1):
        cell = ws.cell(row=row, column=column)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

def auto_width(ws, min_width=10, max_width=45):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = max((len(str(cell.value or "")) for cell in col_cells), default=10)
        ws.column_dimensions[letter].width = min(max(max_len + 2, min_width), max_width)


ws1 = wb.active
ws1.title = "Trade Log"
headers1 = ["Date", "Stock", "Action", "Quantity", "Price", "Trade Value", "Brokerage", "Cash Before", "Cash After", "Notes"]

for column, header in enumerate(headers1, 1):
    ws1.cell(row=1, column=column, value=header)
style_header_row(ws1, 1, len(headers1))

for row_number, trade in enumerate(trade_log, 2):
    date_value = trade["date"].strftime("%Y-%m-%d") if hasattr(trade["date"], "strftime") else str(trade["date"])
    ws1.cell(row=row_number, column=1, value=date_value).font = DATA_FONT
    ws1.cell(row=row_number, column=2, value=trade["symbol"]).font = DATA_FONT
    ws1.cell(row=row_number, column=3, value=trade["action"]).font = DATA_FONT
    ws1.cell(row=row_number, column=4, value=trade["qty"]).font = DATA_FONT
    ws1.cell(row=row_number, column=5, value=trade["price"]).font = DATA_FONT
    ws1.cell(row=row_number, column=5).number_format = NUM_FMT
    ws1.cell(row=row_number, column=6, value=trade["value"]).font = DATA_FONT
    ws1.cell(row=row_number, column=6).number_format = NUM_FMT
    ws1.cell(row=row_number, column=7, value=trade["brokerage"]).font = DATA_FONT
    ws1.cell(row=row_number, column=7).number_format = NUM_FMT
    ws1.cell(row=row_number, column=8, value=trade["cash_before"]).font = DATA_FONT
    ws1.cell(row=row_number, column=8).number_format = NUM_FMT
    ws1.cell(row=row_number, column=9, value=trade["cash_after"]).font = DATA_FONT
    ws1.cell(row=row_number, column=9).number_format = NUM_FMT
    ws1.cell(row=row_number, column=10, value=trade["notes"]).font = SMALL_FONT

    if trade["action"] == "BUY":
        ws1.cell(row=row_number, column=3).fill = BUY_FILL
    elif trade["action"] == "SELL":
        ws1.cell(row=row_number, column=3).fill = SELL_FILL
    elif trade["action"] == "HOLD":
        ws1.cell(row=row_number, column=3).fill = HOLD_FILL

    for column in range(1, len(headers1) + 1):
        ws1.cell(row=row_number, column=column).border = THIN_BORDER

auto_width(ws1)
ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers1))}{len(trade_log) + 1}"
ws1.freeze_panes = "A2"
print(f"  Sheet 1 (Trade Log): {len(trade_log)} transactions")


ws2 = wb.create_sheet("Rebalance Log")
headers2 = ["Date", "Buy Count", "Sell Count", "Hold Count", "Stocks Bought", "Stocks Sold", "Stocks Held", "Cash at End"]

for column, header in enumerate(headers2, 1):
    ws2.cell(row=1, column=column, value=header)
style_header_row(ws2, 1, len(headers2))

for row_number, period in enumerate(all_periods, 2):
    ws2.cell(row=row_number, column=1, value=period["rb_date"]).font = DATA_FONT
    ws2.cell(row=row_number, column=2, value=len(period["buys"])).font = DATA_FONT
    ws2.cell(row=row_number, column=3, value=len(period["sells"])).font = DATA_FONT
    ws2.cell(row=row_number, column=4, value=len(period["holds"])).font = DATA_FONT
    ws2.cell(row=row_number, column=5, value=", ".join(period["buys"])).font = DATA_FONT
    ws2.cell(row=row_number, column=6, value=", ".join(period["sells"])).font = DATA_FONT
    ws2.cell(row=row_number, column=7, value=", ".join(period["holds"])).font = DATA_FONT
    ws2.cell(row=row_number, column=8, value=round(period["cash_at_end"], 2)).font = DATA_FONT
    ws2.cell(row=row_number, column=8).number_format = NUM_FMT

    ws2.cell(row=row_number, column=5).fill = BUY_FILL
    ws2.cell(row=row_number, column=6).fill = SELL_FILL
    ws2.cell(row=row_number, column=7).fill = HOLD_FILL
    ws2.cell(row=row_number, column=8).fill = CASH_FILL

    for column in range(1, len(headers2) + 1):
        ws2.cell(row=row_number, column=column).border = THIN_BORDER

auto_width(ws2)
ws2.auto_filter.ref = f"A1:{get_column_letter(len(headers2))}{len(all_periods) + 1}"
ws2.freeze_panes = "A2"
print(f"  Sheet 2 (Rebalance Log): {len(all_periods)} rows")

ws3 = wb.create_sheet("Backtest Summary")
ws3.cell(row=1, column=1, value="Starting Capital").font = Font(name="Arial", bold=True, size=11)
ws3.cell(row=1, column=2, value=STARTING_CAPITAL)
ws3.cell(row=1, column=2).font = INPUT_FONT
ws3.cell(row=1, column=2).fill = YELLOW_FILL
ws3.cell(row=1, column=2).number_format = "#,##0"
ws3.cell(row=1, column=2).border = THIN_BORDER

if not daily_df.empty:
    total_days = len(daily_df)
    start_date = daily_df["date"].iloc[0]
    end_date = daily_df["date"].iloc[-1]
    years = (end_date - start_date).days / 365.25
    final_value = daily_df["total_value"].iloc[-1]
    total_return = (final_value - STARTING_CAPITAL) / STARTING_CAPITAL
    cagr = (final_value / STARTING_CAPITAL) ** (1 / years) - 1 if years > 0 else 0
    max_drawdown = daily_df["drawdown"].min()
    daily_returns = daily_df["daily_return"].dropna()
    sharpe = (daily_returns.mean() / daily_returns.std()) * np.sqrt(252) if daily_returns.std() > 0 else 0
    annualized_volatility = daily_returns.std() * np.sqrt(252)
    total_brokerage = sum(trade["brokerage"] for trade in trade_log)
    total_trades = len([trade for trade in trade_log if trade["action"] in ("BUY", "SELL")])

    labels_vals = [
        ("Strategy", "High Beta Top 20, No Stop-Loss"),
        ("Period", f"{start_date.strftime('%Y-%m-%d')} to {end_date.strftime('%Y-%m-%d')}"),
        ("Total Trading Days", total_days),
        ("Final Portfolio Value", round(final_value, 2)),
        ("Total Return", total_return),
        ("CAGR", cagr),
        ("Max Drawdown", max_drawdown),
        ("Annualized Volatility", annualized_volatility),
        ("Sharpe Ratio", round(sharpe, 4)),
        ("Avg Daily Return", daily_returns.mean()),
        ("Total Rebalances", len(all_periods)),
        ("Total Trades Executed", total_trades),
        ("Total Brokerage Paid", round(total_brokerage, 2)),
        ("Final Cash Balance", round(cash, 2)),
    ]

    pct_indices = {4, 5, 6, 7, 9}

    index_path = BASE / "nifty_largemidcap_250_data.csv"
    if index_path.exists():
        index_df = pd.read_csv(index_path)
        index_df["date_time"] = pd.to_datetime(index_df["date_time"]).dt.date
        index_df = index_df.set_index("date_time").sort_index()
        mask = (index_df.index >= start_date) & (index_df.index <= end_date)
        period_index = index_df[mask].copy()
        if not period_index.empty:
            start_index_val = float(period_index["close"].iloc[0])
            end_index_val = float(period_index["close"].iloc[-1])
            index_total_return = (end_index_val - start_index_val) / start_index_val
            index_cagr = (end_index_val / start_index_val) ** (1 / years) - 1 if years > 0 else 0
            period_index["peak"] = period_index["close"].cummax()
            period_index["drawdown"] = (period_index["close"] - period_index["peak"]) / period_index["peak"]
            index_max_drawdown = period_index["drawdown"].min()

            labels_vals.extend([
                ("Index Total Return", index_total_return),
                ("Index CAGR", index_cagr),
                ("Index Max Drawdown", index_max_drawdown),
            ])
            idx_start = len(labels_vals) - 3
            pct_indices.update({idx_start, idx_start + 1, idx_start + 2})

    money_indices = {3, 12, 13}
    for idx, (label, value) in enumerate(labels_vals):
        row_number = 3 + idx
        ws3.cell(row=row_number, column=1, value=label).font = Font(name="Arial", bold=True, size=10)
        ws3.cell(row=row_number, column=1).border = THIN_BORDER
        ws3.cell(row=row_number, column=2, value=value).font = DATA_FONT
        ws3.cell(row=row_number, column=2).border = THIN_BORDER

        if idx in pct_indices:
            ws3.cell(row=row_number, column=2).number_format = PCT_FMT
        elif idx in money_indices:
            ws3.cell(row=row_number, column=2).number_format = NUM_FMT

    stock_trades = {}
    for trade in trade_log:
        stock_trades.setdefault(trade["symbol"], []).append(trade)

    stats_start = 3 + len(labels_vals) + 2
    ws3.cell(row=stats_start, column=1, value="Stock Performance Breakdown").font = Font(name="Arial", bold=True, size=12)
    stats_header_row = stats_start + 1
    stats_headers = [
        "Symbol",
        "Total Buys",
        "Total Sells",
        "Total Qty Traded",
        "Total Brokerage",
        "Net Realized PnL",
    ]
    for column, header in enumerate(stats_headers, 1):
        ws3.cell(row=stats_header_row, column=column, value=header)
    style_header_row(ws3, stats_header_row, len(stats_headers))

    stock_stats = []
    for symbol in sorted(stock_trades.keys()):
        trades = stock_trades[symbol]
        buys_count = sum(1 for trade in trades if trade["action"] == "BUY")
        sells_count = sum(1 for trade in trades if trade["action"] == "SELL")
        total_qty = sum(trade["qty"] for trade in trades if trade["action"] != "HOLD")
        total_brokerage = sum(trade["brokerage"] for trade in trades)
        buy_value = sum(trade["value"] for trade in trades if trade["action"] == "BUY")
        sell_value = sum(trade["value"] for trade in trades if trade["action"] == "SELL")
        remaining_qty = portfolio[symbol]["qty"] if symbol in portfolio else 0
        remaining_value = remaining_qty * portfolio[symbol]["last_price"] if symbol in portfolio else 0
        net_pnl = sell_value + remaining_value - buy_value - total_brokerage
        # net_pnl = sell_value - buy_value - total_brokerage
        stock_stats.append((symbol, buys_count, sells_count, total_qty, total_brokerage, net_pnl))

    stock_stats.sort(key=lambda item: -item[5])
    for idx, (symbol, buys_count, sells_count, total_qty, brokerage, net_pnl) in enumerate(stock_stats):
        row_number = stats_header_row + 1 + idx
        ws3.cell(row=row_number, column=1, value=symbol).font = DATA_FONT
        ws3.cell(row=row_number, column=2, value=buys_count).font = DATA_FONT
        ws3.cell(row=row_number, column=3, value=sells_count).font = DATA_FONT
        ws3.cell(row=row_number, column=4, value=total_qty).font = DATA_FONT
        ws3.cell(row=row_number, column=5, value=round(brokerage, 2)).font = DATA_FONT
        ws3.cell(row=row_number, column=5).number_format = NUM_FMT
        ws3.cell(row=row_number, column=6, value=round(net_pnl, 2)).font = DATA_FONT
        ws3.cell(row=row_number, column=6).number_format = NUM_FMT
        for column in range(1, 7):
            ws3.cell(row=row_number, column=column).border = THIN_BORDER

auto_width(ws3)
ws3.freeze_panes = "A3"
print("  Sheet 3 (Backtest Summary): done")


ws4 = wb.create_sheet("Stock Timeline")
dates_list = [period["rb_date"] for period in all_periods]
all_symbols = sorted({trade["symbol"] for trade in trade_log})

ws4.cell(row=1, column=1, value="Symbol")
for column, date_value in enumerate(dates_list, 2):
    ws4.cell(row=1, column=column, value=date_value)
    ws4.cell(row=1, column=column).alignment = Alignment(text_rotation=90, horizontal="center")
style_header_row(ws4, 1, len(dates_list) + 1)

timeline_lookup = {}
for trade in trade_log:
    date_value = trade["date"].strftime("%Y-%m-%d") if hasattr(trade["date"], "strftime") else str(trade["date"])
    timeline_lookup.setdefault((date_value, trade["symbol"]), []).append(trade)

for row_number, symbol in enumerate(all_symbols, 2):
    ws4.cell(row=row_number, column=1, value=symbol).font = Font(name="Arial", bold=True, size=8)
    ws4.cell(row=row_number, column=1).border = THIN_BORDER

    for column, date_value in enumerate(dates_list, 2):
        trades = timeline_lookup.get((date_value, symbol), [])
        if not trades:
            ws4.cell(row=row_number, column=column, value="").border = THIN_BORDER
            continue

        parts = []
        fill = None
        for trade in trades:
            if trade["action"] == "BUY":
                parts.append(f"BUY {trade['qty']}@{trade['price']:.0f}")
                fill = BUY_FILL
            elif trade["action"] == "HOLD":
                parts.append(f"HOLD {trade['qty']}@{trade['price']:.0f}")
                if fill is None:
                    fill = HOLD_FILL
            elif trade["action"] == "SELL":
                parts.append(f"SELL {trade['qty']}@{trade['price']:.0f}")
                fill = SELL_FILL

        cell = ws4.cell(row=row_number, column=column, value=" | ".join(parts))
        cell.font = SMALL_FONT
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")
        if fill:
            cell.fill = fill

ws4.freeze_panes = "B2"
ws4.column_dimensions["A"].width = 14
for column in range(2, len(dates_list) + 2):
    ws4.column_dimensions[get_column_letter(column)].width = 24
print(f"  Sheet 4 (Stock Timeline): {len(all_symbols)} stocks x {len(dates_list)} dates")


ws5 = wb.create_sheet("Daily Portfolio")
headers5 = ["Date", "Stock Value", "Cash", "Total Value", "Daily Return", "Drawdown", "# Stocks"]
for column, header in enumerate(headers5, 1):
    ws5.cell(row=1, column=column, value=header)
style_header_row(ws5, 1, len(headers5))

if not daily_df.empty:
    for idx, daily_row in daily_df.iterrows():
        row_number = 2 + idx
        ws5.cell(row=row_number, column=1, value=daily_row["date"].strftime("%Y-%m-%d")).font = DATA_FONT
        ws5.cell(row=row_number, column=2, value=round(daily_row["stock_value"], 2)).font = DATA_FONT
        ws5.cell(row=row_number, column=2).number_format = NUM_FMT
        ws5.cell(row=row_number, column=3, value=round(daily_row["cash"], 2)).font = DATA_FONT
        ws5.cell(row=row_number, column=3).number_format = NUM_FMT
        ws5.cell(row=row_number, column=4, value=round(daily_row["total_value"], 2)).font = DATA_FONT
        ws5.cell(row=row_number, column=4).number_format = NUM_FMT

        daily_return = daily_row.get("daily_return", 0)
        ws5.cell(row=row_number, column=5, value=round(daily_return, 6) if pd.notna(daily_return) else 0).font = DATA_FONT
        ws5.cell(row=row_number, column=5).number_format = PCT_FMT

        drawdown = daily_row.get("drawdown", 0)
        ws5.cell(row=row_number, column=6, value=round(drawdown, 6) if pd.notna(drawdown) else 0).font = DATA_FONT
        ws5.cell(row=row_number, column=6).number_format = PCT_FMT

        ws5.cell(row=row_number, column=7, value=int(daily_row["num_stocks"])).font = DATA_FONT

        for column in range(1, len(headers5) + 1):
            ws5.cell(row=row_number, column=column).border = THIN_BORDER

auto_width(ws5)
ws5.freeze_panes = "A2"
print(f"  Sheet 5 (Daily Portfolio): {len(daily_df)} rows")

wb.save(OUTPUT_FILE)
print(f"\nSaved: {OUTPUT_FILE.name}")
print(f"complete in {time.time() - start_time:.2f} seconds, ")
